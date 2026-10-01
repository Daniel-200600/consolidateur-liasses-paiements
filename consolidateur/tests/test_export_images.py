"""
tests/test_export_images.py

Fond de feuille dans les classeurs exportés (et absence de logo).

Le fond de feuille n'est pas géré par openpyxl : il est injecté directement
dans le paquet .xlsx (cf. core/export_background.py). Ces tests vérifient que
l'injection produit un fichier toujours valide, sur toutes les feuilles, et
qu'elle reste sans effet sur les formules.
"""

import io
import zipfile
from pathlib import Path

import openpyxl
import pytest

from core.export import build_full_workbook, build_resume_workbook
from core.export_background import apply_background_to_bytes
from core.kpis import compute_kpis, retrait_gab_kpi
from core.parsing import Institution, parse_workbook

FIXTURES = Path(__file__).parent / "fixtures_fictives"


@pytest.fixture
def institutions():
    return [
        Institution(name="BanqueA", file_name="BanqueA.xlsx", sheets=parse_workbook(FIXTURES / "BanqueA.xlsx")),
        Institution(name="OperateurB", file_name="OperateurB.xlsx", sheets=parse_workbook(FIXTURES / "OperateurB.xlsx")),
    ]


def _to_bytes(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def fond(tmp_path):
    """Image de fond quelconque : le contenu n'est pas interprété, seulement copié."""
    chemin = tmp_path / "fond.png"
    chemin.write_bytes(bytes(64))
    return chemin


# --------------------------------------------------------------------- #
# Fond de feuille
# --------------------------------------------------------------------- #

def test_background_is_applied_to_every_sheet(institutions, fond):
    data = apply_background_to_bytes(_to_bytes(build_full_workbook(institutions)), fond)
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        sheets = [n for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
        assert sheets
        for sheet in sheets:
            assert "<picture " in z.read(sheet).decode("utf-8"), f"{sheet} sans fond"
        assert any("media/background.png" in n for n in z.namelist())


def test_workbook_remains_valid_after_injection(institutions, fond):
    """Le fichier retouché doit rester lisible : une injection XML mal formée
    produirait un classeur qu'Excel refuserait d'ouvrir."""
    kpis = compute_kpis(institutions)
    gab = retrait_gab_kpi(institutions)
    data = apply_background_to_bytes(_to_bytes(build_resume_workbook(institutions, kpis, gab)), fond)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Résumé"]


def test_injection_preserves_formulas(institutions, fond):
    kpis = compute_kpis(institutions)
    gab = retrait_gab_kpi(institutions)
    avant = _to_bytes(build_resume_workbook(institutions, kpis, gab))
    apres = apply_background_to_bytes(avant, fond)

    def formules(blob):
        wb = openpyxl.load_workbook(io.BytesIO(blob))
        ws = wb["Résumé"]
        return [c.value for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]

    assert formules(avant) == formules(apres)
    assert len(formules(apres)) > 0


def test_injection_is_idempotent(institutions, fond):
    """Appliquer deux fois le fond ne doit pas empiler deux balises."""
    kpis = compute_kpis(institutions)
    gab = retrait_gab_kpi(institutions)
    once = apply_background_to_bytes(_to_bytes(build_resume_workbook(institutions, kpis, gab)), fond)
    twice = apply_background_to_bytes(once, fond)
    with zipfile.ZipFile(io.BytesIO(twice)) as z:
        xml = z.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert xml.count("<picture ") == 1


def test_missing_image_is_a_no_op(institutions, tmp_path):
    """Si l'image venait à manquer, l'export doit rester exploitable plutôt
    que d'échouer."""
    kpis = compute_kpis(institutions)
    gab = retrait_gab_kpi(institutions)
    data = _to_bytes(build_resume_workbook(institutions, kpis, gab))
    assert apply_background_to_bytes(data, tmp_path / "absente.png") == data


# --------------------------------------------------------------------- #
# Logo imprimable
# --------------------------------------------------------------------- #

def test_logo_beac_retire_du_resume_et_du_sommaire(institutions):
    """Retiré sur demande explicite de l'utilisateur — ne doit plus jamais
    apparaître, ni dans le résumé ni dans le fichier complet."""
    kpis = compute_kpis(institutions)
    gab = retrait_gab_kpi(institutions)
    resume = build_resume_workbook(institutions, kpis, gab)
    assert len(resume["Résumé"]._images) == 0

    full = build_full_workbook(institutions)
    assert len(full["Sommaire"]._images) == 0


# --------------------------------------------------------------------- #
# Analyse qualitative en fin de chaque onglet Total par pays
#
# Demandée explicitement : un résumé complet (dominance + couverture +
# anomalies) à la fin de chaque onglet Total par pays du fichier complet —
# jamais sur l'onglet Total CEMAC (tous pays confondus), qui n'a pas cette
# analyse.
# --------------------------------------------------------------------- #

def _institutions_avec_pays():
    return [
        Institution(name="BanqueA", file_name="BanqueA.xlsx", sheets=parse_workbook(FIXTURES / "BanqueA.xlsx"), pays="CF"),
        Institution(name="OperateurB", file_name="OperateurB.xlsx", sheets=parse_workbook(FIXTURES / "OperateurB.xlsx"), pays="CF"),
    ]


def test_analyse_qualitative_presente_sur_total_pays(institutions):
    from core.export import build_full_workbook
    wb = build_full_workbook(_institutions_avec_pays())
    nom = next(n for n in wb.sheetnames if n.startswith("Mobile Money1-Total-"))
    ws = wb[nom]
    textes = {c.value for row in ws.iter_rows() for c in row if c.value}
    assert any("ANALYSE QUALITATIVE" in str(t) for t in textes)
    assert any("dominant" in str(t).lower() for t in textes)
    assert any("Couverture" in str(t) for t in textes)


def test_analyse_qualitative_absente_du_total_cemac(institutions):
    """Le total CEMAC (tous pays) n'a pas cette analyse — elle n'a de sens
    que par pays, pas pour un ensemble hétérogène de plusieurs pays."""
    from core.export import build_full_workbook
    wb = build_full_workbook(_institutions_avec_pays())
    ws = wb["Mobile Money1-Total"]
    textes = {c.value for row in ws.iter_rows() for c in row if c.value}
    assert not any("ANALYSE QUALITATIVE" in str(t) for t in textes)


def test_anomalie_connue_apparait_dans_lanalyse_qualitative(institutions):
    """Cas réel déjà validé : OperateurB porte une anomalie connue sur
    MobileMoney2 (saut mensuel invraisemblable) — doit apparaître dans
    l'analyse qualitative du total pays correspondant."""
    from core.export import build_full_workbook
    wb = build_full_workbook(_institutions_avec_pays())
    ws = wb["MobileMoney2-Total-CF"]
    textes = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any("⚠️" in t and "OperateurB" in t for t in textes)


def test_recalcul_reel_avec_analyse_qualitative(institutions, tmp_path):
    from core.export import build_full_workbook
    from libreoffice_recalc import recalculer

    chemin = tmp_path / "complet.xlsx"
    build_full_workbook(_institutions_avec_pays()).save(chemin)
    rapport = recalculer(chemin)
    assert rapport["status"] == "success"
    assert rapport["total_errors"] == 0


# --------------------------------------------------------------------- #
# Rapprochements entre établissements différents, dans le Sommaire
# --------------------------------------------------------------------- #

def test_rapprochement_visible_dans_le_sommaire_avec_lien():
    from core.export import build_full_workbook
    from core.parsing import Institution, OngletNonReconnu

    grid_mtn = [["Prêts Credits octroyées aux particuliers en"], [None],
               ["Type de prêt (nom du service)"]]
    grid_ft = [["Prêts Credits octroyées aux particuiers en"], [None],
              ["Type de prêt (nom du service)"], ["Financement fonds de roulement"],
              ["Crédit à la consommation"], ["Crédit équipement"], ["Crédit scolaire"]]
    mtn = Institution("MTN", "mtn.xlsx", {}, non_reconnus=[OngletNonReconnu("Prets-Credits 2025", grid_mtn)])
    ft = Institution("FIRSTTRUST", "ft.xlsx", {}, non_reconnus=[OngletNonReconnu("PretsCredits 2025", grid_ft)])

    wb = build_full_workbook([mtn, ft])
    ws = wb["Sommaire"]
    textes = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert any("RAPPROCHEMENTS POSSIBLES" in t for t in textes)
    assert "MTN" in textes and "FIRSTTRUST" in textes


def test_pas_de_bloc_rapprochement_si_rien_a_rapprocher(institutions):
    """Sur les fichiers de référence (BanqueA/OperateurB, tous deux entièrement
    reconnus), aucun bloc de rapprochement ne doit apparaître."""
    from core.export import build_full_workbook
    wb = build_full_workbook(institutions)
    ws = wb["Sommaire"]
    textes = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
    assert not any("RAPPROCHEMENTS POSSIBLES" in t for t in textes)


# --------------------------------------------------------------------- #
# Feuilles sans information exploitable (Sheet1 vide, etc.) — ignorées
# --------------------------------------------------------------------- #

def test_feuille_vide_est_totalement_ignoree(tmp_path):
    from openpyxl import Workbook
    from core.parsing import analyser_fichier

    wb = Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Sheet1")  # complètement vide, nom générique par défaut
    chemin = tmp_path / "avec_feuille_vide.xlsx"
    wb.save(chemin)

    sheets, non_reconnus = analyser_fichier(chemin)
    assert sheets == {}
    assert non_reconnus == []  # ni consolidée, ni listée hors gabarit


def test_feuille_nom_generique_avec_vraies_donnees_nest_pas_ignoree():
    from core.parsing import _onglet_sans_information_exploitable
    grid = [["Nombre de contrats souscrits", 4200], ["Primes collectées", 89000000]]
    assert not _onglet_sans_information_exploitable("Sheet1", grid)


def test_freeze_panes_najamais_un_pan_vertical_profond(institutions):
    """Cas réel signalé : le résumé figeait 61 lignes sur un document de
    plusieurs milliers, remplissant tout l'écran visible d'une zone gelée
    et rendant le début du document impossible à atteindre."""
    from core.export import build_resume_workbook
    from core.kpis import compute_kpis, retrait_gab_kpi
    wb = build_resume_workbook(institutions, compute_kpis(institutions), retrait_gab_kpi(institutions))
    assert wb["Résumé"].freeze_panes == "B1"


def test_pas_de_ligne_etablissements_consolides(institutions):
    """Retirée sur demande explicite — rien concernant les fichiers
    individuellement dans les exports."""
    from core.export import build_full_workbook, build_resume_workbook
    from core.kpis import compute_kpis, retrait_gab_kpi

    for wb, feuille in [(build_full_workbook(institutions), "Sommaire"),
                        (build_resume_workbook(institutions, compute_kpis(institutions),
                                               retrait_gab_kpi(institutions)), "Résumé")]:
        textes = {c.value for row in wb[feuille].iter_rows(max_row=5) for c in row if c.value}
        assert not any("Établissements consolidés" in str(t) for t in textes)


