"""
tests/libreoffice_recalc.py

Recalcul réel d'un classeur par LibreOffice, pour les tests d'export.

openpyxl écrit les formules sans valeur en cache : seul un vrai tableur peut
dire si elles produisent #REF!, #VALUE!… LibreOffice (sans interface) ouvre
le fichier, calcule toutes les formules et le réenregistre au même endroit ;
les tests relisent ensuite les valeurs calculées avec openpyxl
(`data_only=True`).

LibreOffice est cherché dans le PATH (Linux, macOS, CI) puis dans les
dossiers d'installation Windows. S'il est introuvable, le test appelant est
ignoré (skip) avec ce motif plutôt que d'échouer pour une raison étrangère au
code testé.
"""

from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from collections import Counter
from pathlib import Path

import openpyxl
import pytest

ERREURS_EXCEL = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!")

# Profil LibreOffice dédié aux tests : évite les conflits avec une instance
# ouverte par l'utilisateur et l'assistant de premier démarrage.
_PROFIL = Path(tempfile.gettempdir()) / "consolidateur-tests-libreoffice"


def trouver_soffice() -> str | None:
    for nom in ("soffice", "libreoffice"):
        chemin = shutil.which(nom)
        if chemin:
            return chemin
    for racine in (os.environ.get("ProgramFiles"), os.environ.get("ProgramFiles(x86)")):
        if racine:
            candidat = Path(racine) / "LibreOffice" / "program" / "soffice.exe"
            if candidat.exists():
                return str(candidat)
    return None


def _compter_formules(chemin: Path) -> int:
    wb = openpyxl.load_workbook(chemin)
    return sum(
        1
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    )


def recalculer(chemin: str | Path, timeout: int = 300) -> dict:
    """Recalcule `chemin` en place et renvoie le rapport d'erreurs de formule.

    Rapport : status (« success » ou « errors_found »), total_errors,
    total_formulas, error_summary (nombre d'occurrences par type d'erreur).
    """
    soffice = trouver_soffice()
    if soffice is None:
        pytest.skip("LibreOffice introuvable : recalcul réel des formules impossible")

    chemin = Path(chemin)
    total_formules = _compter_formules(chemin)

    with tempfile.TemporaryDirectory() as sortie:
        subprocess.run(
            [
                soffice,
                f"-env:UserInstallation={_PROFIL.as_uri()}",
                "--headless", "--norestore",
                "--convert-to", "xlsx:Calc MS Excel 2007 XML",
                "--outdir", sortie, str(chemin),
            ],
            capture_output=True, timeout=timeout, check=False,
        )
        produit = Path(sortie) / chemin.name
        if not produit.exists():
            raise RuntimeError(f"LibreOffice n'a pas produit de fichier recalculé pour {chemin.name}")
        shutil.copyfile(produit, chemin)

    erreurs = Counter(
        cell.value
        for ws in openpyxl.load_workbook(chemin, data_only=True).worksheets
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value in ERREURS_EXCEL
    )
    total = sum(erreurs.values())
    return {
        "status": "success" if total == 0 else "errors_found",
        "total_errors": total,
        "total_formulas": total_formules,
        "error_summary": dict(erreurs),
    }
