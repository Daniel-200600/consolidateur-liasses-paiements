"""
tests/test_resume_exhaustivite.py

Le résumé exporté doit rendre compte de TOUT le gabarit : chaque feuille,
chaque titre de tableau, chaque ligne — renseignée ou non — et chaque mesure,
avec des valeurs exactes et sans ligne ambiguë.

Ces tests figent un audit qui avait révélé quatre manques :
- la feuille PrêtsCredits, absente du récapitulatif parce qu'elle ne comporte
  aucune ligne prédéfinie ;
- 13 titres de tableaux sur 26, ceux dont aucune ligne n'était renseignée ;
- 194 lignes de contenu sur 249, toutes celles laissées vides ;
- deux lignes « Total B » indiscernables dans Cartes 5, le titre de tableau
  « TRANSACTIONS DE VOS PORTEURS » y ouvrant deux blocs distincts.
"""

from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from core.export import build_resume_workbook
from core.indicators import discover_indicators, indicator_totals
from core.kpis import compute_kpis, retrait_gab_kpi
from core.parsing import KNOWN_SHEETS, Institution, parse_workbook
from libreoffice_recalc import recalculer

FIXTURES = Path(__file__).parent / "fixtures_fictives"


@pytest.fixture(scope="module")
def institutions():
    return [
        Institution(name="BanqueA", file_name="BanqueA.xlsx", sheets=parse_workbook(FIXTURES / "BanqueA.xlsx")),
        Institution(name="OperateurB", file_name="OperateurB.xlsx", sheets=parse_workbook(FIXTURES / "OperateurB.xlsx")),
    ]


@pytest.fixture(scope="module")
def resume(institutions, tmp_path_factory):
    kpis = compute_kpis(institutions)
    gab = retrait_gab_kpi(institutions)
    path = tmp_path_factory.mktemp("resume") / "resume.xlsx"
    build_resume_workbook(institutions, kpis, gab).save(path)
    # openpyxl n'écrit aucune valeur en cache : sans recalcul, les colonnes
    # calculées se relisent comme des chaînes « =SUM(...) ».
    recalculer(path)
    return openpyxl.load_workbook(path, data_only=True)["Résumé"]


@pytest.fixture(scope="module")
def textes(resume):
    """Tous les textes présents dans la feuille."""
    out = set()
    for row in resume.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip():
                out.add(cell.value.strip())
    return out


@pytest.fixture(scope="module")
def lignes_detail(resume):
    """Bloc 3 indexé par (feuille, tableau, indicateur, mesure)."""
    out = {}
    for r in range(1, resume.max_row + 1):
        libelle = resume.cell(row=r, column=3).value
        if isinstance(libelle, str) and libelle[:1] in "●○":
            cle = (
                resume.cell(row=r, column=1).value,
                resume.cell(row=r, column=2).value,
                libelle[2:],
                resume.cell(row=r, column=4).value,
            )
            out[cle] = tuple(resume.cell(row=r, column=c).value or 0 for c in (5, 6, 7))
    return out


def test_toutes_les_feuilles_sont_citees(textes):
    manquantes = [s for s in KNOWN_SHEETS if s not in textes]
    assert not manquantes, f"Feuilles absentes du résumé : {manquantes}"


def test_feuille_sans_ligne_predefinie_est_citee(textes):
    """PrêtsCredits est une grille libre : elle doit tout de même figurer."""
    assert "PrêtsCredits" in textes


def test_tous_les_titres_de_tableaux_sont_cites(institutions, textes):
    rubriques = {(i.sheet, i.rubric) for i in discover_indicators(institutions) if i.rubric}
    manquants = [r for r in rubriques if r[1] not in textes]
    assert not manquants, f"Titres de tableaux absents : {manquants}"


def test_toutes_les_lignes_de_contenu_sont_reprises(institutions, textes):
    indicateurs = discover_indicators(institutions)
    manquantes = [
        i for i in indicateurs
        if f"● {i.display_name}" not in textes and f"○ {i.display_name}" not in textes
    ]
    assert not manquantes, f"{len(manquantes)} ligne(s) absentes, ex. {manquantes[:3]}"


def test_les_lignes_vides_sont_reprises_et_marquees(institutions, textes):
    """Une ligne non renseignée est une information : elle doit figurer,
    signalée par ○."""
    vides = [i for i in discover_indicators(institutions) if not i.has_data]
    assert len(vides) > 150
    assert all(f"○ {i.display_name}" in textes for i in vides)


def test_chaque_ligne_est_identifiable_sans_ambiguite(institutions, lignes_detail):
    """Deux lignes du résumé ne doivent jamais partager la même identité
    (feuille, tableau, indicateur, mesure) tout en portant des valeurs
    différentes."""
    attendu = sum(len(i.measures) for i in discover_indicators(institutions))
    assert len(lignes_detail) == attendu


def test_toutes_les_valeurs_sont_exactes(institutions, lignes_detail):
    ecarts = []
    for ind in discover_indicators(institutions):
        for mesure in ind.measures:
            cle = (ind.sheet, ind.rubric or "—", ind.display_name, mesure or "Valeur")
            par_etab, total = indicator_totals(institutions, ind, mesure)
            attendu = (par_etab[0][1], par_etab[1][1], total)
            if lignes_detail.get(cle) != attendu:
                ecarts.append((cle, lignes_detail.get(cle), attendu))
    assert not ecarts, f"{len(ecarts)} valeur(s) incorrectes, ex. {ecarts[:2]}"


def test_les_blocs_repetes_sont_numerotes(institutions):
    """Cartes 5 comporte deux blocs « TRANSACTIONS DE VOS PORTEURS » ; le
    second doit être distingué, sans quoi ses lignes seraient confondues."""
    rubriques = {i.rubric for i in discover_indicators(institutions) if i.sheet == "Cartes 5"}
    assert "TRANSACTIONS DE VOS PORTEURS" in rubriques
    assert "TRANSACTIONS DE VOS PORTEURS (2)" in rubriques


# --------------------------------------------------------------------- #
# Énumérations de réseaux monétiques
# --------------------------------------------------------------------- #

RESEAUX = ["Privative", "GIMAC", "VISA", "MASTERCARD", "UnionPay", "American Express"]


@pytest.mark.parametrize("reseau", RESEAUX)
def test_chaque_reseau_monetique_est_restitue(institutions, textes, reseau):
    """Les réseaux non renseignés (VISA dans Cartes 3, MASTERCARD dans
    Cartes 1…) étaient pris pour des titres de tableau et disparaissaient du
    résumé, faute de valeur chez l'un des établissements."""
    assert any(reseau in t for t in textes), f"Réseau absent du résumé : {reseau}"


def test_les_deux_blocs_de_cartes1_sont_distingues(institutions):
    """Cartes 1 liste les mêmes réseaux deux fois : une fois pour les
    transactions, une fois pour la fraude. Chaque occurrence doit rester
    identifiable et porter ses propres mesures."""
    inds = [i for i in discover_indicators(institutions) if i.sheet == "Cartes1"]
    visa = [i for i in inds if i.label.strip().upper() == "VISA"]
    assert len(visa) == 2
    assert len({i.display_name for i in visa}) == 2
    mesures = {tuple(i.measures) for i in visa}
    assert len(mesures) == 2, "Les deux blocs doivent avoir des mesures distinctes"


def test_bloc_fraude_sur_les_cartes_est_present(institutions, textes):
    assert any("Fraude sur les cartes" in t for t in textes)


def test_un_reseau_nest_jamais_un_titre_de_tableau(institutions):
    rubriques = {i.rubric for i in discover_indicators(institutions)}
    for reseau in ("VISA", "MASTERCARD", "UnionPay"):
        assert not any(r.strip().upper().startswith(reseau.upper()) for r in rubriques), \
            f"{reseau} ne doit pas être un titre de tableau"


# --------------------------------------------------------------------- #
# Navigation dans la feuille
# --------------------------------------------------------------------- #

def test_le_bloc_de_detail_est_filtrable(institutions, tmp_path):
    """Le bloc de détail dépasse 500 lignes : un filtre est indispensable
    pour n'afficher qu'une feuille, qu'un tableau ou que les lignes
    renseignées."""
    kpis = compute_kpis(institutions)
    gab = retrait_gab_kpi(institutions)
    chemin = tmp_path / "resume.xlsx"
    build_resume_workbook(institutions, kpis, gab).save(chemin)
    ws = openpyxl.load_workbook(chemin)["Résumé"]

    assert ws.auto_filter.ref, "Aucun filtre automatique posé"
    debut, fin = ws.auto_filter.ref.split(":")
    assert debut.startswith("A")
    # Le filtre doit couvrir toutes les colonnes, jusqu'au total marché.
    assert fin[0] == get_column_letter(5 + len(institutions))
    # et l'en-tête doit rester visible au défilement
    assert ws.freeze_panes is not None


# --------------------------------------------------------------------- #
# Notes et définitions du gabarit (bloc 5)
#
# Ces notes étaient déjà copiées telles quelles dans le fichier complet
# (recopie brute des onglets sources) mais absentes du résumé : quelqu'un
# ne consultant que ce document n'avait aucun moyen de savoir, par exemple,
# ce que recouvre précisément un « porteur actif ».
# --------------------------------------------------------------------- #

def test_les_notes_du_gabarit_sont_toutes_reprises(institutions, textes):
    from core.indicators import extraire_notes
    notes = extraire_notes(institutions)
    assert len(notes) == 16
    for _, texte in notes:
        assert any(texte in t or t in texte for t in textes), f"Note absente du résumé : {texte[:60]}"


def test_les_notes_sont_dedupliquees_entre_etablissements(institutions):
    from core.indicators import extraire_notes
    notes = extraire_notes(institutions)
    assert len(notes) == len(set(notes))


def test_note_specifique_porteur_actif_est_lisible(institutions, textes):
    assert any("effectué au moins une transaction" in t for t in textes)


# --------------------------------------------------------------------- #
# Onglets hors gabarit standard (bloc 6)
#
# Un onglet qu'aucun thème commun ne recouvre (mise en page propre à un
# établissement, ou onglet composite non encore segmenté) ne doit jamais
# disparaître des documents livrés : au minimum, sa présence est tracée
# dans le résumé, et son contenu intégral reste accessible dans le fichier
# complet (cf. tests/test_export.py pour la vérification du contenu).
# --------------------------------------------------------------------- #



def test_pas_de_bloc_hors_gabarit_si_tout_est_reconnu(institutions, tmp_path):
    """Sur les fichiers de référence (100% reconnus), le bloc ne doit pas
    apparaître du tout — pas de bruit inutile."""
    from core.export import build_resume_workbook
    from core.kpis import compute_kpis, retrait_gab_kpi
    import openpyxl

    chemin = tmp_path / "resume.xlsx"
    build_resume_workbook(institutions, compute_kpis(institutions), retrait_gab_kpi(institutions)).save(chemin)
    ws = openpyxl.load_workbook(chemin)["Résumé"]
    textes_resume = {str(c.value) for row in ws.iter_rows() for c in row if c.value}
    assert not any("HORS GABARIT" in t for t in textes_resume)


# --------------------------------------------------------------------- #
# Bloc 4 (points de vigilance) : le libellé français doit couvrir tous
# les types d'alerte, y compris les incohérences de totaux hiérarchiques
# — trouvé en réauditant un export réel à grande échelle (47
# établissements) où 16 incohérences de ce type existaient.
# --------------------------------------------------------------------- #

