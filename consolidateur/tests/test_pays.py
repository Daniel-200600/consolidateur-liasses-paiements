"""
tests/test_pays.py

Regroupement par pays de la CEMAC (Cameroun, Centrafrique, Congo, Gabon,
Guinée Équatoriale, Tchad) : détection automatique depuis le nom de
fichier, totaux par pays et vue CEMAC d'ensemble dans les deux exports.

Le codage du pays n'est pas uniforme d'un fichier à l'autre — constaté
concrètement sur les tout premiers fichiers reçus : un établissement
centrafricain a nommé son fichier « BanqueA_CA-… », un autre « OperateurB_RCA_… »,
pour le même pays. La détection doit couvrir ces variantes, mais rester
toujours modifiable : un pays mal détecté fausserait un total pays sans
qu'aucune erreur ne soit signalée.
"""

from pathlib import Path

import openpyxl
import pytest

from core.export import build_full_workbook, build_resume_workbook
from core.kpis import compute_kpis, retrait_gab_kpi
from core.parsing import Institution, parse_workbook
from core.pays import NOM_PAYS, PAYS_CEMAC, detecter_pays
from libreoffice_recalc import recalculer

FIXTURES = Path(__file__).parent / "fixtures_fictives"


# --------------------------------------------------------------------- #
# Détection depuis le nom de fichier
# --------------------------------------------------------------------- #

@pytest.mark.parametrize("nom, attendu", [
    ("BanqueA_CA-Liasse_Stats_Paiements_2025__1_.xlsx", "CF"),
    ("OperateurB_RCA_Liasse_Stats_Paiements_2025_1.xlsx", "CF"),
    ("ETABA-CM_Liasse_Stats_Paiements_2025_2026.xlsx", "CM"),
    ("Banque_du_Gabon_Liasse_2025.xlsx", "GA"),
    ("TCHAD_BANK_Stats_2025.xlsx", "TD"),
    ("GQ-BANK_2025.xlsx", "GQ"),
    ("CONGO-CG_Liasse_2025.xlsx", "CG"),
])
def test_detection_sur_noms_reels_et_plausibles(nom, attendu):
    assert detecter_pays(nom) == attendu


def test_meme_pays_deux_codages_differents():
    """Cas réel : deux établissements du même pays, deux conventions de
    nommage différentes pour ce pays."""
    assert detecter_pays("BanqueA_CA-Liasse_Stats_Paiements_2025__1_.xlsx") == \
           detecter_pays("OperateurB_RCA_Liasse_Stats_Paiements_2025_1.xlsx") == "CF"


def test_code_court_najamais_de_faux_positif_en_sous_chaine():
    """« CA » ne doit jamais être détecté à l'intérieur d'un mot — sans quoi
    « CartesBank » ou « CashApp » seraient à tort rattachés à la
    Centrafrique."""
    assert detecter_pays("CartesBank_Statistiques_2025.xlsx") is None
    assert detecter_pays("CashApp_Rapport_2025.xlsx") is None


def test_code_composite_non_pays_reste_non_detecte():
    """« ACA » (un nom de prestataire, sur le modèle d'un cas réel) ne doit correspondre à
    aucun code pays, même s'il contient des lettres qui s'en approchent."""
    assert detecter_pays("ETATS_TRANSFERT_ACA_BANK_2025.xlsx") is None


def test_nom_de_fichier_sans_pays_renvoie_none():
    assert detecter_pays("Liasse_Stats_Paiements_2025.xlsx") is None


def test_six_pays_cemac_couverts():
    assert set(PAYS_CEMAC.keys()) == {"CM", "CF", "CG", "GA", "GQ", "TD"}
    assert all(v["nom"] for v in PAYS_CEMAC.values())


# --------------------------------------------------------------------- #
# Totaux par pays dans le fichier complet
# --------------------------------------------------------------------- #

@pytest.fixture
def deux_pays():
    F = FIXTURES
    banque_a = Institution(name="BanqueA", file_name="b.xlsx", sheets=parse_workbook(F / "BanqueA.xlsx"), pays="CF")
    operateur_b = Institution(name="OperateurB", file_name="o.xlsx", sheets=parse_workbook(F / "OperateurB.xlsx"), pays="CF")
    banque_a_cm = Institution(name="BanqueA-CM", file_name="b2.xlsx", sheets=parse_workbook(F / "BanqueA.xlsx"), pays="CM")
    return [banque_a, operateur_b, banque_a_cm]


def _recalc(chemin: Path) -> dict:
    return recalculer(chemin)


def test_un_onglet_total_par_pays_present(deux_pays, tmp_path):
    chemin = tmp_path / "complet.xlsx"
    wb = build_full_workbook(deux_pays)
    wb.save(chemin)
    assert "Mobile Money1-Total-CF" in wb.sheetnames
    assert "Mobile Money1-Total-CM" in wb.sheetnames
    assert "Mobile Money1-Total" in wb.sheetnames  # total CEMAC, toujours présent aussi


def test_total_pays_narrondit_rien_recalcul_reel(deux_pays, tmp_path):
    chemin = tmp_path / "complet.xlsx"
    build_full_workbook(deux_pays).save(chemin)
    rapport = _recalc(chemin)
    assert rapport["status"] == "success"
    assert rapport["total_errors"] == 0


def test_total_pays_cf_egale_les_valeurs_de_reference(deux_pays, tmp_path):
    """BanqueA + OperateurB, tous deux en Centrafrique : le total pays CF doit
    retomber exactement sur les valeurs déjà validées de longue date."""
    chemin = tmp_path / "complet.xlsx"
    build_full_workbook(deux_pays).save(chemin)
    _recalc(chemin)
    ws = openpyxl.load_workbook(chemin, data_only=True)["Mobile Money1-Total-CF"]
    assert ws["M7"].value == 1232671


def test_total_cemac_reste_juste_avec_plusieurs_pays(deux_pays, tmp_path):
    """Le total CEMAC (tous pays) doit rester identique, que les
    établissements soient répartis sur un ou plusieurs pays."""
    chemin = tmp_path / "complet.xlsx"
    build_full_workbook(deux_pays).save(chemin)
    _recalc(chemin)
    ws = openpyxl.load_workbook(chemin, data_only=True)["Mobile Money1-Total"]
    assert ws["M7"].value == 1232671  # BanqueA-CM n'apporte rien sur ce thème


def test_etablissement_sans_pays_najamais_de_total_pays(tmp_path):
    """Un établissement dont le pays n'a pas été détecté ni renseigné ne
    doit jamais faire planter la génération, ni produire un onglet pays
    fantôme."""
    F = FIXTURES
    banque_a = Institution(name="BanqueA", file_name="b.xlsx", sheets=parse_workbook(F / "BanqueA.xlsx"), pays=None)
    chemin = tmp_path / "complet.xlsx"
    wb = build_full_workbook([banque_a])
    wb.save(chemin)
    assert not any("-Total-" in n for n in wb.sheetnames)
    assert "Cartes1-Total" in wb.sheetnames


def test_sommaire_reference_les_totaux_pays(deux_pays):
    wb = build_full_workbook(deux_pays)
    som = wb["Sommaire"]
    textes = {str(c.value) for row in som.iter_rows() for c in row if c.value}
    assert any("TOTAUX PAR PAYS" in t for t in textes)
    assert "Cameroun" in textes
    assert "République Centrafricaine" in textes


# --------------------------------------------------------------------- #
# Vue par pays dans le résumé
# --------------------------------------------------------------------- #

def test_bloc_par_pays_present_dans_le_resume(deux_pays, tmp_path):
    chemin = tmp_path / "resume.xlsx"
    kpis, gab = compute_kpis(deux_pays), retrait_gab_kpi(deux_pays)
    build_resume_workbook(deux_pays, kpis, gab).save(chemin)
    _recalc(chemin)
    ws = openpyxl.load_workbook(chemin, data_only=True)["Résumé"]
    textes = {str(c.value) for row in ws.iter_rows() for c in row if c.value}
    assert any("PAR PAYS" in t for t in textes)


def test_colonne_cemac_du_resume_egale_le_total_marche(deux_pays, tmp_path):
    """Le total CEMAC du tableau « Montant total par pays » doit toujours
    égaler la somme des totaux par pays qui le précèdent — jamais un
    chiffre recalculé séparément qui pourrait diverger."""
    from core.montant_total import montant_total_argent

    chemin = tmp_path / "resume.xlsx"
    kpis, gab = compute_kpis(deux_pays), retrait_gab_kpi(deux_pays)
    build_resume_workbook(deux_pays, kpis, gab).save(chemin)
    _recalc(chemin)
    ws = openpyxl.load_workbook(chemin, data_only=True)["Résumé"]

    r_titre = next(r for r in range(1, ws.max_row + 1)
                  if ws.cell(row=r, column=1).value and "MONTANT TOTAL PAR PAYS" in str(ws.cell(row=r, column=1).value))
    lignes_pays = {}
    r = r_titre + 2  # saute le titre et l'en-tête de colonnes
    while ws.cell(row=r, column=1).value and ws.cell(row=r, column=1).value != "CEMAC (total)":
        lignes_pays[ws.cell(row=r, column=1).value] = ws.cell(row=r, column=2).value
        r += 1
    total_cemac_affiche = ws.cell(row=r, column=2).value

    assert set(lignes_pays.keys()) == {"Cameroun", "République Centrafricaine"}
    assert sum(lignes_pays.values()) == pytest.approx(total_cemac_affiche)
    assert total_cemac_affiche == pytest.approx(montant_total_argent(deux_pays))


def test_pas_de_bloc_pays_si_aucun_pays_renseigne(tmp_path):
    F = FIXTURES
    banque_a = Institution(name="BanqueA", file_name="b.xlsx", sheets=parse_workbook(F / "BanqueA.xlsx"), pays=None)
    chemin = tmp_path / "resume.xlsx"
    kpis, gab = compute_kpis([banque_a]), retrait_gab_kpi([banque_a])
    build_resume_workbook([banque_a], kpis, gab).save(chemin)
    ws = openpyxl.load_workbook(chemin)["Résumé"]
    textes = {str(c.value) for row in ws.iter_rows() for c in row if c.value}
    assert not any("PAR PAYS" in t for t in textes)


# --------------------------------------------------------------------- #
# Le tableau « Montant total par pays » ne vit plus que dans le résumé —
# retiré du Sommaire du fichier complet pour ne pas y faire doublon avec
# les onglets Total par pays (qui, eux, restent).
# --------------------------------------------------------------------- #

def test_montant_total_par_pays_absent_du_fichier_complet(deux_pays, tmp_path):
    chemin = tmp_path / "complet.xlsx"
    build_full_workbook(deux_pays).save(chemin)
    rapport = _recalc(chemin)
    assert rapport["status"] == "success" and rapport["total_errors"] == 0

    ws = openpyxl.load_workbook(chemin, data_only=True)["Sommaire"]
    textes = {str(c.value) for row in ws.iter_rows() for c in row if c.value}
    assert not any("MONTANT TOTAL PAR PAYS" in t for t in textes)
    # mais les onglets Total par pays (formules), eux, restent bien présents
    assert "Mobile Money1-Total-CF" in build_full_workbook(deux_pays).sheetnames


def test_montant_total_par_pays_reste_dans_le_resume(deux_pays, tmp_path):
    from core.montant_total import montant_total_argent

    chemin = tmp_path / "resume.xlsx"
    kpis, gab = compute_kpis(deux_pays), retrait_gab_kpi(deux_pays)
    build_resume_workbook(deux_pays, kpis, gab).save(chemin)
    _recalc(chemin)

    ws = openpyxl.load_workbook(chemin, data_only=True)["Résumé"]
    textes_debut = {ws.cell(row=r, column=1).value for r in range(1, 40)}
    assert any(t and "MONTANT TOTAL PAR PAYS" in str(t) for t in textes_debut)

    valeurs = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(1, 40)}
    assert valeurs.get("Cameroun") == pytest.approx(montant_total_argent([deux_pays[2]]))
    assert valeurs.get("République Centrafricaine") == pytest.approx(
        montant_total_argent([deux_pays[0], deux_pays[1]]))

