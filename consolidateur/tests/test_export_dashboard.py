"""
tests/test_export_dashboard.py

Feuille « Tableau de bord » du fichier complet.

Deux exigences, dont la seconde corrige un premier essai insuffisant :
- chaque chiffre est une FORMULE pointant vers l'onglet Total correspondant
  (jamais une valeur recopiée), de sorte que la feuille se recalcule avec le
  classeur ;
- TOUS les indicateurs y figurent — le premier essai n'en retenait que 13,
  choisis à la main, ce qui laissait le reste invisible malgré une
  consolidation par ailleurs correcte.
"""

from pathlib import Path

import openpyxl
import pytest

from core.export import build_full_workbook
from core.export_dashboard import _adresse_total
from core.indicators import discover_indicators, indicator_totals
from core.parsing import Institution, KNOWN_SHEETS, parse_workbook
from libreoffice_recalc import recalculer

FIXTURES = Path(__file__).parent / "fixtures_fictives"


@pytest.fixture(scope="module")
def institutions():
    return [
        Institution(name="BanqueA", file_name="BanqueA.xlsx", sheets=parse_workbook(FIXTURES / "BanqueA.xlsx")),
        Institution(name="OperateurB", file_name="OperateurB.xlsx", sheets=parse_workbook(FIXTURES / "OperateurB.xlsx")),
    ]


@pytest.fixture(scope="module")
def classeur(institutions):
    return build_full_workbook(institutions)


@pytest.fixture(scope="module")
def feuille_recalculee(institutions, tmp_path_factory):
    chemin = tmp_path_factory.mktemp("dash") / "complet.xlsx"
    build_full_workbook(institutions).save(chemin)
    recalculer(chemin)
    return openpyxl.load_workbook(chemin, data_only=True)["Tableau de bord"]


def _lignes_detail(ws):
    """Bloc de détail indexé par (feuille, tableau, indicateur, mesure)."""
    out = {}
    for r in range(1, ws.max_row + 1):
        libelle = ws.cell(row=r, column=3).value
        if isinstance(libelle, str) and libelle[:1] in "●○":
            cle = (ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value,
                  libelle[2:], ws.cell(row=r, column=4).value)
            out[cle] = ws.cell(row=r, column=5).value
    return out


def test_la_feuille_existe_et_est_la_premiere(classeur):
    assert "Tableau de bord" in classeur.sheetnames
    assert classeur.sheetnames[0] == "Tableau de bord"


def test_tous_les_indicateurs_sont_presents_sans_exception(institutions, feuille_recalculee):
    """Rien n'est laissé de côté : le nombre de lignes de détail doit
    correspondre exactement au nombre de couples (indicateur, mesure)
    détectés dans les fichiers, comme dans la section « Tous les
    indicateurs » de l'application et le bloc 3 du résumé."""
    indicateurs = discover_indicators(institutions)
    attendu = sum(len(i.measures) for i in indicateurs)
    lignes = _lignes_detail(feuille_recalculee)
    assert len(lignes) == attendu


def test_toutes_les_feuilles_du_gabarit_sont_couvertes(institutions, feuille_recalculee):
    lignes = _lignes_detail(feuille_recalculee)
    feuilles_presentes = {cle[0] for cle in lignes}
    indicateurs = discover_indicators(institutions)
    feuilles_avec_indicateur = {i.sheet for i in indicateurs}
    assert feuilles_avec_indicateur <= feuilles_presentes


def test_toutes_les_valeurs_sont_exactes(institutions, feuille_recalculee):
    indicateurs = discover_indicators(institutions)
    lignes = _lignes_detail(feuille_recalculee)
    ecarts = []
    for ind in indicateurs:
        for mesure in ind.measures:
            cle = (ind.sheet, ind.rubric or "—", ind.display_name, mesure or "Valeur")
            _, total = indicator_totals(institutions, ind, mesure)
            if lignes.get(cle) != total:
                ecarts.append((cle, lignes.get(cle), total))
    assert not ecarts, f"{len(ecarts)} valeur(s) incorrecte(s), ex. {ecarts[:3]}"


def test_les_valeurs_sont_des_formules_pas_des_valeurs_recopiees(classeur):
    """Le cœur de l'exigence : aucune valeur figée."""
    ws = classeur["Tableau de bord"]
    colonne_valeur = [
        ws.cell(row=r, column=5).value
        for r in range(1, ws.max_row + 1)
        if isinstance(ws.cell(row=r, column=3).value, str)
        and ws.cell(row=r, column=3).value[:1] in "●○"
    ]
    assert len(colonne_valeur) > 400
    non_formule = [v for v in colonne_valeur if not (isinstance(v, str) and v.startswith("=N("))]
    assert not non_formule, f"Valeurs non branchées par formule : {non_formule[:5]}"


def test_le_dashboard_reflete_une_modification_de_la_source(institutions, tmp_path):
    """Preuve que le tableau de bord est vivant : on modifie une valeur dans
    un onglet établissement, et le chiffre correspondant doit suivre après
    recalcul — quelle que soit la ligne concernée."""
    chemin = tmp_path / "complet.xlsx"
    build_full_workbook(institutions).save(chemin)

    wb = openpyxl.load_workbook(chemin)
    onglet = next(n for n in wb.sheetnames if n.startswith("Cartes1") and n.endswith("BanqueA"))
    wb[onglet]["B3"] = 11386 + 1000       # « Clients porteurs de cartes »
    wb.save(chemin)

    recalculer(chemin)
    ws = openpyxl.load_workbook(chemin, data_only=True)["Tableau de bord"]
    lignes = _lignes_detail(ws)
    valeur = next(v for (sheet, _, lib, _), v in lignes.items()
                 if sheet == "Cartes1" and "clients" in lib.lower())
    assert valeur == 12386


def test_les_compteurs_de_couverture_sont_structurels_non_selectifs(institutions, feuille_recalculee):
    """La bande du haut ne met aucun indicateur métier en avant : seuls des
    compteurs de couverture (nb feuilles, nb indicateurs, taux de
    remplissage) y figurent."""
    ws = feuille_recalculee
    valeurs = {ws.cell(row=6, column=c).value: ws.cell(row=7, column=c).value for c in [1, 3, 5, 7]}
    assert valeurs.get("ÉTABLISSEMENTS") == len(institutions)
    assert valeurs.get("INDICATEURS SUIVIS") == len(discover_indicators(institutions))


def test_le_bloc_de_detail_est_filtrable(classeur):
    ws = classeur["Tableau de bord"]
    assert ws.auto_filter.ref
    assert ws.freeze_panes is not None


def test_aucune_erreur_de_formule(institutions, tmp_path):
    chemin = tmp_path / "complet.xlsx"
    build_full_workbook(institutions).save(chemin)
    rapport = recalculer(chemin)
    assert rapport["status"] == "success"
    assert rapport["total_errors"] == 0


def test_le_lien_retour_sommaire_est_valide(classeur):
    ws = classeur["Tableau de bord"]
    liens = [c for row in ws.iter_rows() for c in row if c.hyperlink and c.hyperlink.location]
    assert liens
    assert all("Sommaire" in c.hyperlink.location for c in liens)


def test_adresse_total_pointe_la_bonne_cellule(institutions):
    assert _adresse_total(institutions, "Cartes 5", 12, "Valeur") == "C13"
    adresse = _adresse_total(institutions, "Mobile Money1", 6, "")
    assert adresse and adresse.endswith("7")


def test_indicateur_absent_est_ignore_sans_erreur():
    banque_a = Institution(name="BanqueA", file_name="b.xlsx", sheets=parse_workbook(FIXTURES / "BanqueA.xlsx"))
    wb = build_full_workbook([banque_a])
    assert "Tableau de bord" in wb.sheetnames


def test_une_cartographie_validee_est_prise_en_compte(institutions):
    """La feuille doit honorer une cartographie validée : une ligne
    requalifiée en titre par le modèle local ne doit pas y apparaître comme
    une donnée — cohérence avec la section « Tous les indicateurs » de
    l'application."""
    from core.cartographie import Cartographie, EntreeCartographie
    from core.llm import ROLE_TITRE

    carto = Cartographie()
    carto.entrees[("Cartes1", 19)] = EntreeCartographie(
        sheet="Cartes1", row=19, libelle="MASTERCARD", role=ROLE_TITRE)
    carto.entrees[("Cartes1", 31)] = EntreeCartographie(
        sheet="Cartes1", row=31, libelle="MASTERCARD", role=ROLE_TITRE)

    wb = build_full_workbook(institutions, cartographie=carto)
    ws = wb["Tableau de bord"]
    # MASTERCARD apparaît légitimement dans d'autres feuilles (Cartes 3) : la
    # vérification porte donc sur (feuille, libellé), pas sur le libellé seul.
    trouve = any(
        ws.cell(row=r, column=1).value == "Cartes1"
        and isinstance(ws.cell(row=r, column=3).value, str)
        and "MASTERCARD" in ws.cell(row=r, column=3).value
        for r in range(1, ws.max_row + 1)
    )
    assert not trouve, "MASTERCARD de Cartes1 aurait dû être retiré par la cartographie"
