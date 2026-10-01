"""
core/export_style.py

Charte graphique des exports Excel : couleurs, polices, bordures et
utilitaires de mise en forme, repris du thème de l'application
(vert émeraude / or, cf. .streamlit/config.toml).

Isolé du module d'export pour que la logique de calcul (quelles valeurs,
quelles formules) reste lisible indépendamment de l'habillage.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- Palette ---------------------------------------------------------------
EMERAUDE = "0B6E4F"
EMERAUDE_FONCE = "0B4A34"
EMERAUDE_CLAIR = "E3F3EC"
OR = "C9A227"
# Nuance distincte du total CEMAC (OR) : les totaux par pays doivent se
# repérer d'un coup d'œil dans une liste d'onglets qui en compte beaucoup.
OR_PAYS = "8A6D1D"
OR_FONCE = "7A5E12"
CREME = "FBF3D9"
BLANC = "FFFFFF"
GRIS_TEXTE = "16261D"
GRIS_LEGER = "F5F5F3"
ROUGE = "B03A2E"
# Onglets hors gabarit standard (cf. Étape 1 suite) : ni un établissement, ni
# un total, ni le sommaire — une teinte à part pour les repérer d'un coup
# d'œil dans la barre d'onglets.
AMBRE = "C87137"

# Police : Arial, exigée pour tout livrable (cf. skill xlsx).
POLICE = "Arial"

# --- Polices ---------------------------------------------------------------
F_TITRE = Font(name=POLICE, size=14, bold=True, color=BLANC)
F_SOUS_TITRE = Font(name=POLICE, size=10, italic=True, color=EMERAUDE_FONCE)
F_ENTETE = Font(name=POLICE, size=10, bold=True, color=BLANC)
F_RUBRIQUE = Font(name=POLICE, size=10, bold=True, color=EMERAUDE_FONCE)
F_SOUS_SECTION = Font(name=POLICE, size=10, bold=True, color=OR_FONCE)
F_LIBELLE = Font(name=POLICE, size=10, color=GRIS_TEXTE)
F_DONNEE = Font(name=POLICE, size=10, color=GRIS_TEXTE)
F_TOTAL = Font(name=POLICE, size=10, bold=True, color=EMERAUDE_FONCE)
F_NOTE = Font(name=POLICE, size=8, italic=True, color=OR_FONCE)
F_ERREUR = Font(name=POLICE, size=9, italic=True, color=ROUGE)
F_LIEN = Font(name=POLICE, size=10, color=EMERAUDE, underline="single")
# Lien posé sur le bandeau de titre (fond émeraude) : il doit rester lisible,
# d'où une teinte claire plutôt que le vert du lien courant.
F_LIEN_SOMMAIRE = Font(name=POLICE, size=10, bold=True, color=CREME, underline="single")

# --- Remplissages ----------------------------------------------------------
FILL_TITRE = PatternFill("solid", fgColor=EMERAUDE)
FILL_ENTETE = PatternFill("solid", fgColor=EMERAUDE)
FILL_RUBRIQUE = PatternFill("solid", fgColor=EMERAUDE_CLAIR)
FILL_SOUS_SECTION = PatternFill("solid", fgColor=CREME)
FILL_TOTAL = PatternFill("solid", fgColor=CREME)
FILL_ALTERNE = PatternFill("solid", fgColor=GRIS_LEGER)

# --- Bordures --------------------------------------------------------------
_fin = Side(style="thin", color="D5D5D0")
_moyen = Side(style="medium", color=EMERAUDE)
_or = Side(style="thin", color=OR)

B_GRILLE = Border(left=_fin, right=_fin, top=_fin, bottom=_fin)
B_ENTETE = Border(left=_fin, right=_fin, top=_moyen, bottom=_moyen)
B_TOTAL = Border(left=_or, right=_or, top=_or, bottom=_or)

# --- Formats de nombre -----------------------------------------------------
# Un zéro s'affiche « - » : sur ces liasses, l'immense majorité des lignes du
# gabarit ne sont pas renseignées, et une colonne de zéros rend le tableau
# illisible.
FMT_NOMBRE = '#,##0;-#,##0;"-"'
FMT_MOIS = "mmm-yy"

ALIGN_LIBELLE = Alignment(horizontal="left", vertical="center", wrap_text=False)
ALIGN_NOMBRE = Alignment(horizontal="right", vertical="center")
ALIGN_ENTETE = Alignment(horizontal="center", vertical="center", wrap_text=True)


def ajuster_largeurs(ws, largeur_libelle: int = 46, largeur_donnee: int = 16,
                     max_col: int | None = None) -> None:
    """Largeur généreuse pour la colonne des libellés, uniforme ensuite."""
    ws.column_dimensions["A"].width = largeur_libelle
    dernier = max_col or ws.max_column
    for c in range(2, dernier + 1):
        ws.column_dimensions[get_column_letter(c)].width = largeur_donnee


def couleur_onglet(ws, couleur: str) -> None:
    ws.sheet_properties.tabColor = couleur
